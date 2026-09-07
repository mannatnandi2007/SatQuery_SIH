"""
SatQuery AI — Multi-Format Report Generator
Assembles query results and telemetry into downloadable reports across 4 formats:
1. JSON (.json)
2. Markdown (.md)
3. PDF (.pdf) via ReportLab
4. Microsoft Word (.docx) via python-docx
"""

import os
import json
import uuid
from datetime import datetime, timezone
from typing import Dict, List, Optional

# ReportLab imports for PDF generation
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, HRFlowable, KeepTogether
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

# python-docx imports for Word document generation
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


# Directories
REPORTS_DIR = os.path.join(os.path.dirname(__file__), "reports")
OVERLAYS_DIR = os.path.join(os.path.dirname(__file__), "static", "overlays")
os.makedirs(REPORTS_DIR, exist_ok=True)


# ─── Markdown Generator ──────────────────────────────────────────

def generate_markdown(report_data: Dict, output_path: str):
    """Generate clean, structured GitHub-flavored Markdown inspection report."""
    q_data = report_data.get("query", {})
    r_data = report_data.get("result", {})
    trace_data = report_data.get("execution_trace", [])
    analysis = r_data.get("detailed_analysis") or {}

    md = []
    md.append(f"# SatQuery AI — Remote Sensing Inspection Report")
    md.append(f"**Report ID:** `{report_data.get('report_id')}`  ")
    md.append(f"**Timestamp:** `{report_data.get('generated_at')}`  ")
    md.append(f"**System Version:** `{report_data.get('system', {}).get('version', 'SatQuery AI v2.4')}`  ")
    md.append(f"**Specialist Model:** `{report_data.get('system', {}).get('model', 'RS-VLM')}`\n")
    md.append("---\n")

    md.append("## 1. Query & Ingestion Metadata\n")
    md.append(f"- **Visual Query:** {q_data.get('text', 'N/A')}")
    images = q_data.get("images", [])
    md.append(f"- **Acquisition Count:** {len(images)} scene(s)")
    for i, img_name in enumerate(images, 1):
        md.append(f"  - **Scene {i}:** `{img_name}`")
    md.append(f"- **Evidence Grounding:** `{r_data.get('evidence', {}).get('type', 'none')}`")
    conf = r_data.get("confidence", {})
    md.append(f"- **Confidence Score:** **{conf.get('label', 'N/A')}** ({conf.get('score', 0.0):.2f})\n")

    md.append("## 2. Executive Findings\n")
    md.append(f"> {r_data.get('answer', 'No summary available.')}\n")

    if analysis:
        md.append("## 3. Multi-Angle Satellite Telemetry\n")
        md.append(f"### Scene Overview\n{analysis.get('scene_overview', 'N/A')}\n")
        md.append(f"### Land Cover & Transition (LULC)\n{analysis.get('land_cover', 'N/A')}\n")

        key_objs = analysis.get("key_objects", [])
        if key_objs:
            md.append("### Key Spatial Features & Infrastructure")
            for obj in key_objs:
                md.append(f"- {obj}")
            md.append("")

        md.append(f"### Spatial & Structural Configuration\n{analysis.get('spatial_patterns', 'N/A')}\n")
        md.append(f"### Spectral Indicators & Sensor Observations\n{analysis.get('spectral_observations', 'N/A')}\n")
        md.append(f"### Operational & Environmental Risk Assessment\n{analysis.get('potential_concerns', 'N/A')}\n")

    overlay_file = r_data.get("evidence", {}).get("overlay_file")
    if overlay_file:
        md.append("## 4. Visual Evidence Artifact\n")
        md.append(f"![Visual Evidence Overlay](/static/overlays/{overlay_file})\n")

    if trace_data:
        md.append("## 5. Execution Trace & Latency Ledger\n")
        md.append("| Stage | Status | Details | Duration (ms) |")
        md.append("| :--- | :--- | :--- | :--- |")
        for step in trace_data:
            md.append(f"| {step.get('stage')} | {step.get('status').upper()} | {step.get('detail')} | {step.get('duration_ms', 0):.1f} |")
        md.append(f"\n**Total Duration:** `{report_data.get('total_duration_ms', 0):.1f} ms`\n")

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(md))


# ─── PDF Generator (ReportLab) ───────────────────────────────────

def generate_pdf(report_data: Dict, output_path: str, overlay_path: Optional[str] = None):
    """Generate a high-precision, publication-grade PDF report using ReportLab."""
    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()

    # Custom styles adhering to Cobalt theme
    c_primary = colors.HexColor("#0284c7")
    c_dark = colors.HexColor("#0f172a")
    c_muted = colors.HexColor("#64748b")
    c_light_bg = colors.HexColor("#f8fafc")
    c_border = colors.HexColor("#cbd5e1")

    title_style = ParagraphStyle(
        "DocTitle",
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=c_dark
    )
    subtitle_style = ParagraphStyle(
        "DocSubTitle",
        fontName="Helvetica",
        fontSize=10,
        leading=13,
        textColor=c_primary
    )
    h2_style = ParagraphStyle(
        "SectionH2",
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=c_dark,
        spaceBefore=12,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        "Body",
        fontName="Helvetica",
        fontSize=9.5,
        leading=14,
        textColor=colors.HexColor("#334155")
    )
    callout_style = ParagraphStyle(
        "Callout",
        fontName="Helvetica",
        fontSize=10,
        leading=15,
        textColor=c_dark
    )
    table_cell = ParagraphStyle(
        "TableCell",
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=c_dark
    )
    table_cell_bold = ParagraphStyle(
        "TableCellBold",
        fontName="Helvetica-Bold",
        fontSize=8.5,
        leading=11,
        textColor=c_dark
    )

    story = []

    # Title & Header
    story.append(Paragraph("SATQUERY AI — EARTH OBSERVATION REPORT", subtitle_style))
    story.append(Paragraph("Remote Sensing Inspection & Telemetry", title_style))
    story.append(Spacer(1, 8))
    story.append(HRFlowable(width="100%", thickness=1.5, color=c_primary, spaceBefore=2, spaceAfter=10))

    # Metadata Table
    q_data = report_data.get("query", {})
    r_data = report_data.get("result", {})
    conf = r_data.get("confidence", {})

    meta_content = [
        [
            Paragraph("<b>Report ID:</b>", table_cell_bold), Paragraph(str(report_data.get("report_id")), table_cell),
            Paragraph("<b>Generated:</b>", table_cell_bold), Paragraph(str(report_data.get("generated_at"))[:19], table_cell)
        ],
        [
            Paragraph("<b>Confidence:</b>", table_cell_bold), Paragraph(f"<b>{conf.get('label')}</b> ({conf.get('score', 0):.2f})", table_cell),
            Paragraph("<b>Imagery Count:</b>", table_cell_bold), Paragraph(f"{len(q_data.get('images', []))} scene(s)", table_cell)
        ],
        [
            Paragraph("<b>Visual Query:</b>", table_cell_bold), Paragraph(str(q_data.get("text")), table_cell),
            Paragraph("<b>Evidence Type:</b>", table_cell_bold), Paragraph(str(r_data.get("evidence", {}).get("type")), table_cell)
        ]
    ]
    meta_table = Table(meta_content, colWidths=[80, 190, 80, 190])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), c_light_bg),
        ('BOX', (0, 0), (-1, -1), 0.5, c_border),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 12))

    # Executive Summary Box
    story.append(Paragraph("Executive Summary & Findings", h2_style))
    summary_text = r_data.get("answer", "No analysis recorded.")
    summary_box = Table([[Paragraph(summary_text, callout_style)]], colWidths=[540])
    summary_box.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#f0f9ff")),
        ('BOX', (0, 0), (-1, -1), 1, c_primary),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('LEFTPADDING', (0, 0), (-1, -1), 12),
        ('RIGHTPADDING', (0, 0), (-1, -1), 12),
    ]))
    story.append(summary_box)
    story.append(Spacer(1, 10))

    # Visual Evidence Image (if available)
    if overlay_path and os.path.exists(overlay_path):
        story.append(Paragraph("Visual Evidence Grounding", h2_style))
        try:
            # Resize image to fit nicely within printable width (540pt)
            story.append(RLImage(overlay_path, width=520, height=260))
            story.append(Spacer(1, 10))
        except Exception as e:
            print(f"[PDF] Could not embed overlay: {e}")

    # Detailed Analysis Telemetry Grid
    analysis = r_data.get("detailed_analysis") or {}
    if analysis:
        story.append(Paragraph("Multi-Angle Satellite Telemetry", h2_style))
        telemetry_rows = [
            [Paragraph("<b>Scene Overview</b>", table_cell_bold), Paragraph(analysis.get("scene_overview", "N/A"), table_cell)],
            [Paragraph("<b>Land Cover (LULC)</b>", table_cell_bold), Paragraph(analysis.get("land_cover", "N/A"), table_cell)],
            [Paragraph("<b>Spatial Patterns</b>", table_cell_bold), Paragraph(analysis.get("spatial_patterns", "N/A"), table_cell)],
            [Paragraph("<b>Spectral Indicators</b>", table_cell_bold), Paragraph(analysis.get("spectral_observations", "N/A"), table_cell)],
            [Paragraph("<b>Risk & Exposure</b>", table_cell_bold), Paragraph(analysis.get("potential_concerns", "N/A"), table_cell)],
        ]
        key_items = analysis.get("key_objects", [])
        if key_items:
            obj_str = "<br/>• ".join([""] + key_items)[5:]
            telemetry_rows.insert(2, [Paragraph("<b>Key Features</b>", table_cell_bold), Paragraph(obj_str, table_cell)])

        t_table = Table(telemetry_rows, colWidths=[120, 420])
        t_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), c_light_bg),
            ('BOX', (0, 0), (-1, -1), 0.5, c_border),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))
        story.append(t_table)
        story.append(Spacer(1, 10))

    # Execution Trace
    trace_data = report_data.get("execution_trace", [])
    if trace_data:
        story.append(Paragraph("Execution Trace & Telemetry Audit", h2_style))
        trace_rows = [[
            Paragraph("<b>Stage</b>", table_cell_bold),
            Paragraph("<b>Status</b>", table_cell_bold),
            Paragraph("<b>Detail</b>", table_cell_bold),
            Paragraph("<b>Latency</b>", table_cell_bold)
        ]]
        for step in trace_data:
            trace_rows.append([
                Paragraph(step.get("stage", ""), table_cell),
                Paragraph(step.get("status", "").upper(), table_cell),
                Paragraph(step.get("detail", "")[:70], table_cell),
                Paragraph(f"{step.get('duration_ms', 0):.1f}ms", table_cell)
            ])

        tr_table = Table(trace_rows, colWidths=[100, 50, 320, 70])
        tr_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), c_light_bg),
            ('BOX', (0, 0), (-1, -1), 0.5, c_border),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        story.append(tr_table)

    doc.build(story)


# ─── Word Generator (python-docx) ────────────────────────────────

def generate_docx(report_data: Dict, output_path: str, overlay_path: Optional[str] = None):
    """Generate a clean, styled Microsoft Word (.docx) inspection report."""
    doc = docx.Document()

    # Configure 0.6 inch margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.6)
        s.right_margin = Inches(0.6)

    # Document Header
    p_meta = doc.add_paragraph()
    run_sub = p_meta.add_run("SATQUERY AI — EARTH OBSERVATION WORKBENCH\n")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(9.5)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(2, 132, 199)

    run_title = p_meta.add_run("Remote Sensing Inspection Report")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    # Metadata Table
    q_data = report_data.get("query", {})
    r_data = report_data.get("result", {})
    conf = r_data.get("confidence", {})

    table = doc.add_table(rows=3, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(1.2), Inches(2.4), Inches(1.2), Inches(2.4)]
    meta_rows = [
        [("Report ID", True), (str(report_data.get("report_id")), False), ("Timestamp", True), (str(report_data.get("generated_at"))[:19], False)],
        [("Confidence", True), (f"{conf.get('label')} ({conf.get('score', 0):.2f})", False), ("Images", True), (f"{len(q_data.get('images', []))} scene(s)", False)],
        [("Query", True), (str(q_data.get("text")), False), ("Evidence", True), (str(r_data.get("evidence", {}).get("type")), False)]
    ]

    for r_idx, row in enumerate(meta_rows):
        for c_idx, (text, is_bold) in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.width = col_widths[c_idx]
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(2)
            cp.paragraph_format.space_after = Pt(2)
            c_run = cp.add_run(text)
            c_run.font.name = "Arial"
            c_run.font.size = Pt(8.5)
            c_run.font.bold = is_bold
            if is_bold:
                c_run.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Executive Summary Heading & Callout
    h_exec = doc.add_heading("Executive Summary & Findings", level=2)
    h_exec.paragraph_format.space_before = Pt(10)
    h_exec.paragraph_format.space_after = Pt(4)

    p_summary = doc.add_paragraph()
    p_summary.paragraph_format.left_indent = Inches(0.2)
    p_summary.paragraph_format.right_indent = Inches(0.2)
    p_summary.paragraph_format.space_before = Pt(4)
    p_summary.paragraph_format.space_after = Pt(8)
    run_sum = p_summary.add_run(r_data.get("answer", "No analysis recorded."))
    run_sum.font.name = "Arial"
    run_sum.font.size = Pt(10)
    run_sum.font.color.rgb = RGBColor(15, 23, 42)

    # Evidence Image
    if overlay_path and os.path.exists(overlay_path):
        h_ev = doc.add_heading("Visual Evidence Grounding", level=2)
        h_ev.paragraph_format.space_before = Pt(10)
        h_ev.paragraph_format.space_after = Pt(4)
        try:
            doc.add_picture(overlay_path, width=Inches(6.8))
            p_cap = doc.add_paragraph("Figure: Orthorectified satellite scene with detected spatial features and evidence overlay.")
            p_cap.paragraph_format.space_after = Pt(8)
            p_cap.runs[0].font.size = Pt(8)
            p_cap.runs[0].font.italic = True
        except Exception as e:
            print(f"[Word DOCX] Could not embed overlay: {e}")

    # Detailed Telemetry
    analysis = r_data.get("detailed_analysis") or {}
    if analysis:
        h_tel = doc.add_heading("Multi-Angle Satellite Telemetry", level=2)
        h_tel.paragraph_format.space_before = Pt(10)
        h_tel.paragraph_format.space_after = Pt(4)

        t_table = doc.add_table(rows=0, cols=2)
        t_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_widths = [Inches(1.8), Inches(5.4)]

        items = [
            ("Scene Overview", analysis.get("scene_overview", "N/A")),
            ("Land Cover (LULC)", analysis.get("land_cover", "N/A")),
            ("Key Spatial Features", "\n• ".join([""] + analysis.get("key_objects", []))[3:] if analysis.get("key_objects") else "N/A"),
            ("Spatial Configuration", analysis.get("spatial_patterns", "N/A")),
            ("Spectral Indicators", analysis.get("spectral_observations", "N/A")),
            ("Risk Assessment", analysis.get("potential_concerns", "N/A")),
        ]

        for label, val in items:
            row = t_table.add_row()
            c0, c1 = row.cells[0], row.cells[1]
            c0.width, c1.width = t_widths[0], t_widths[1]

            p0 = c0.paragraphs[0]
            r0 = p0.add_run(label)
            r0.font.name = "Arial"
            r0.font.bold = True
            r0.font.size = Pt(8.5)

            p1 = c1.paragraphs[0]
            r1 = p1.add_run(str(val))
            r1.font.name = "Arial"
            r1.font.size = Pt(8.5)

    # Save to disk
    doc.save(output_path)


# ─── Main Report Orchestration ────────────────────────────────────

def generate_report(
    query: str,
    image_filenames: List[str],
    answer: str,
    confidence: Dict,
    trace: List[Dict],
    evidence_type: str,
    overlay_filename: Optional[str] = None,
    detailed_analysis: Optional[Dict] = None,
    model: Optional[str] = None,
    dl_metrics: Optional[Dict] = None
) -> Dict:
    """
    Generate all 4 report formats (JSON, Markdown, PDF, DOCX) and save to disk.

    Returns:
        {"report_id": str, "report_url": str}
    """
    report_id = uuid.uuid4().hex[:16]
    timestamp = datetime.now(timezone.utc).isoformat()

    report_data = {
        "report_id": report_id,
        "generated_at": timestamp,
        "query": {
            "text": query,
            "images": image_filenames,
        },
        "result": {
            "answer": answer,
            "detailed_analysis": detailed_analysis,
            "confidence": confidence,
            "dl_metrics": dl_metrics,
            "evidence": {
                "type": evidence_type,
                "overlay_file": overlay_filename
            }
        },
        "execution_trace": trace,
        "total_duration_ms": sum(t.get("duration_ms", 0) for t in trace),
        "system": {
            "version": "SatQuery AI v2.4 (Cobalt Telemetry)",
            "model": model or "RS-VLM / Change Detection / SAR Fusion",
            "pipeline": "compatibility → routing → specialist → evidence → report"
        }
    }

    # 1. Save JSON Report
    json_path = os.path.join(REPORTS_DIR, f"report_{report_id}.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(report_data, f, indent=2, ensure_ascii=False)

    # 2. Save Markdown Report (.md)
    md_path = os.path.join(REPORTS_DIR, f"report_{report_id}.md")
    try:
        generate_markdown(report_data, md_path)
    except Exception as e:
        print(f"[Report] Markdown export failed: {e}")

    # Determine overlay path if exists
    overlay_path = None
    if overlay_filename:
        cand_path = os.path.join(OVERLAYS_DIR, overlay_filename)
        if os.path.exists(cand_path):
            overlay_path = cand_path

    # 3. Save PDF Report (.pdf)
    pdf_path = os.path.join(REPORTS_DIR, f"report_{report_id}.pdf")
    try:
        generate_pdf(report_data, pdf_path, overlay_path)
    except Exception as e:
        print(f"[Report] PDF export failed: {e}")

    # 4. Save Word Report (.docx)
    docx_path = os.path.join(REPORTS_DIR, f"report_{report_id}.docx")
    try:
        generate_docx(report_data, docx_path, overlay_path)
    except Exception as e:
        print(f"[Report] Word DOCX export failed: {e}")

    return {
        "report_id": report_id,
        "report_url": f"/report/{report_id}"
    }
